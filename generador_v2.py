"""
generador_v2.py  v4.0
Generador de Evaluaciones AMR
UI: historial | chat | opciones
"""

import os, sys, json, threading, random, copy, uuid
from datetime import datetime, date
import tkinter as tk
from tkinter import filedialog, messagebox
import customtkinter as ctk

from config import get_api_key, set_api_key
from gemini_client import GeminiClient
from file_parser import extraer_texto
from generador_app import Generador, LOGO_PATH, AMR_PATH

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

ctk.set_appearance_mode("dark")
ctk.set_default_color_theme("blue")

# ── Paleta oscura ─────────────────────────────────────────────────────────────
BG       = "#141414"
PANEL    = "#1C1C1E"
SIDEBAR  = "#161618"
BORDER   = "#2C2C2E"
MUTED    = "#8E8E93"
TEXT     = "#F2F2F7"
TEXT2    = "#D1D1D6"
ACCENT   = "#0A84FF"
ACCENT_D = "#0070E0"
GREEN    = "#30D158"
GREEN_D  = "#25A244"
PURPLE   = "#BF5AF2"
PURPLE_D = "#9A3FD4"
RED      = "#FF453A"
CHAT_BG  = "#111113"
BUB_U    = "#0A84FF"
BUB_IA   = "#2C2C2E"

FONT = "Segoe UI"

# ── Datos ─────────────────────────────────────────────────────────────────────
CURSOS = [
    "1° Básico", "2° Básico", "3° Básico", "4° Básico",
    "5° Básico", "6° Básico", "7° Básico", "8° Básico",
    "1° Medio",  "2° Medio",  "3° Medio",  "4° Medio",
]

ASIGNATURAS = [
    "Lenguaje", "Matemáticas", "Historia", "Ciencias Naturales",
    "Inglés", "Ed. Física", "Arte", "Música", "Tecnología", "Otra",
]

TIPOS_PREGUNTA = [
    "Selección múltiple",
    "Verdadero / Falso",
    "Completar texto",
    "Desarrollo",
    "Mixta",
]

# ── Persistencia ──────────────────────────────────────────────────────────────
_CONV_DIR = os.path.join(
    os.environ.get("APPDATA", os.path.expanduser("~")),
    "GeneradorAMR", "conversations"
)


def _guardar_conv_disco(conv: dict) -> None:
    os.makedirs(_CONV_DIR, exist_ok=True)
    path = os.path.join(_CONV_DIR, f"{conv['id']}.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(conv, f, ensure_ascii=False, indent=2)


def _cargar_convs_disco() -> list:
    if not os.path.isdir(_CONV_DIR):
        return []
    convs = []
    for fname in os.listdir(_CONV_DIR):
        if fname.endswith(".json"):
            try:
                with open(os.path.join(_CONV_DIR, fname), encoding="utf-8") as f:
                    convs.append(json.load(f))
            except Exception:
                pass
    return sorted(convs, key=lambda c: c.get("date", ""), reverse=True)


# ── Clave de respuestas ───────────────────────────────────────────────────────
def generar_clave_docx(fichas: list, path: str):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("CLAVE DE RESPUESTAS")
    r.bold = True; r.font.size = Pt(14)
    doc.add_paragraph()
    for ficha in fichas:
        p = doc.add_paragraph()
        p.add_run(f"Ficha N° {ficha.get('numero', '')}").bold = True
        t = doc.add_table(rows=1, cols=2)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        for cell, txt in zip(t.rows[0].cells, ["Pregunta", "Respuesta"]):
            cell.text = txt
            cell.paragraphs[0].runs[0].bold = True
        for preg in ficha.get("preguntas", []):
            row = t.add_row().cells
            row[0].text = str(preg.get("numero", ""))
            tipo = preg.get("tipo", "seleccion_multiple")
            if tipo == "desarrollo":
                row[1].text = "Desarrollo libre"
            else:
                row[1].text = preg.get("respuesta_correcta", "")
        doc.add_paragraph()
    doc.save(path)


# ── Versión B ─────────────────────────────────────────────────────────────────
def _hacer_version_b(ficha: dict) -> dict:
    """
    Crea versión B de una ficha:
    - Baraja el orden de las preguntas
    - Dentro de cada pregunta, baraja las alternativas y actualiza respuesta_correcta
    """
    LETRAS = ["A", "B", "C", "D", "E"]

    def extraer_contenido(alt: str) -> str:
        if len(alt) > 1 and alt[1] in (")", " "):
            return alt[2:].strip()
        return alt

    fb = copy.deepcopy(ficha)
    preguntas = fb.get("preguntas", [])
    random.shuffle(preguntas)

    for i, preg in enumerate(preguntas):
        preg["numero"] = i + 1
        alts = preg.get("alternativas", [])
        if not alts:
            continue

        resp_letra = preg.get("respuesta_correcta", "A")

        # Separar contenido de cada alternativa
        contenidos = [extraer_contenido(a) for a in alts]

        # Encontrar contenido de la alternativa correcta
        contenido_correcto = None
        for alt in alts:
            if alt.startswith(f"{resp_letra})") or alt.startswith(f"{resp_letra} "):
                contenido_correcto = extraer_contenido(alt)
                break
        if contenido_correcto is None and contenidos:
            contenido_correcto = contenidos[0]

        # Barajar contenidos
        random.shuffle(contenidos)

        # Reconstruir con nuevas letras y encontrar la correcta
        nuevas_alts = []
        nueva_letra_correcta = LETRAS[0]
        for j, contenido in enumerate(contenidos):
            letra = LETRAS[j]
            nuevas_alts.append(f"{letra}) {contenido}")
            if contenido == contenido_correcto:
                nueva_letra_correcta = letra

        preg["alternativas"] = nuevas_alts
        preg["respuesta_correcta"] = nueva_letra_correcta

    fb["preguntas"] = preguntas
    return fb


# ════════════════════════════════════════════════════════════════════════════
#  App
# ════════════════════════════════════════════════════════════════════════════
class App(ctk.CTk):

    def __init__(self):
        super().__init__()
        self.title("Generador de Evaluaciones  ·  AMR")
        self.geometry("1300x800")
        self.minsize(1020, 660)
        self.configure(fg_color=BG)
        self.tk.call("tk", "scaling", 1.0)
        try:
            self.wm_attributes("-transparentcolor", "")
        except Exception:
            pass
        try:
            icon_path = os.path.join(
                getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__))),
                "icono.ico"
            )
            if os.path.isfile(icon_path):
                self.iconbitmap(icon_path)
        except Exception:
            pass

        self._gemini: GeminiClient | None = None
        self._fichas: list | None         = None
        self._texto_base                  = ""
        self._archivo_nombre              = ""
        self._imagen_bytes: bytes | None  = None
        self._imagen_nombre               = ""
        self._generando                   = False
        self._chat_w                      = 500
        self._resize_job                  = None
        self._pensando_widget             = None

        # Conversación activa
        self._conv_id: str | None    = None
        self._conv_messages: list    = []

        self._verificar_api_key()
        self._build_ui()
        self.bind("<Configure>", self._on_resize)
        self._nueva_conv(silencioso=True)

    # ── API Key ───────────────────────────────────────────────────────────────
    def _verificar_api_key(self):
        key = get_api_key()
        if key:
            self._gemini = GeminiClient(key)
        else:
            self._pedir_api_key()

    def _pedir_api_key(self):
        dlg = ctk.CTkToplevel(self)
        dlg.title("API Key de Gemini")
        dlg.geometry("480x340")
        dlg.resizable(False, False)
        dlg.configure(fg_color=PANEL)
        dlg.grab_set()

        ctk.CTkLabel(dlg, text="Configura tu API Key",
                     text_color=TEXT,
                     font=ctk.CTkFont(FONT, 18, "bold")).pack(
                         padx=32, pady=(32, 4), anchor="w")
        ctk.CTkLabel(dlg, text="Consíguela gratis en aistudio.google.com",
                     text_color=MUTED,
                     font=ctk.CTkFont(FONT, 11)).pack(padx=32, anchor="w")

        entry = ctk.CTkEntry(dlg, width=416, show="•",
                             placeholder_text="Pega tu API Key aquí",
                             font=ctk.CTkFont("Consolas", 11),
                             fg_color=BG, text_color=TEXT,
                             border_color=BORDER, border_width=1,
                             height=44, corner_radius=10)
        entry.pack(padx=32, pady=(24, 6))

        lbl_err = ctk.CTkLabel(dlg, text="", text_color=RED,
                               font=ctk.CTkFont(FONT, 10))
        lbl_err.pack(padx=32, anchor="w")

        btn_ok = ctk.CTkButton(dlg, text="Guardar y continuar",
                               fg_color=ACCENT, hover_color=ACCENT_D,
                               text_color="white",
                               font=ctk.CTkFont(FONT, 12, "bold"),
                               height=44, corner_radius=10,
                               command=lambda: _guardar())
        btn_ok.pack(padx=32, pady=(16, 0), fill="x")

        def _guardar():
            key = entry.get().strip()
            if not key:
                lbl_err.configure(text="Pega tu API Key primero.", text_color=RED)
                return
            lbl_err.configure(text="Verificando…", text_color=MUTED)
            btn_ok.configure(state="disabled", text="Verificando…")
            dlg.update_idletasks()

            def _verify():
                cli = GeminiClient(key)
                ok, err = cli.verificar_api_key()
                dlg.after(0, lambda: _done(cli, ok, err))

            def _done(cli, ok, err):
                if ok:
                    set_api_key(key)
                    self._gemini = cli
                    dlg.destroy()
                else:
                    lbl_err.configure(
                        text=f"Key inválida: {err[:65]}", text_color=RED)
                    btn_ok.configure(state="normal", text="Guardar y continuar")

            threading.Thread(target=_verify, daemon=True).start()

        self.wait_window(dlg)

    # ── Build UI ──────────────────────────────────────────────────────────────
    def _build_ui(self):
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(1, weight=1)
        self._build_header()
        self._build_main()

    def _build_header(self):
        hdr = ctk.CTkFrame(self, fg_color=PANEL, corner_radius=0, height=62)
        hdr.grid(row=0, column=0, sticky="ew")
        hdr.grid_propagate(False)
        hdr.grid_columnconfigure(1, weight=1)

        if AMR_PATH:
            try:
                from PIL import Image
                img = ctk.CTkImage(Image.open(AMR_PATH).convert("RGBA"),
                                   size=(36, 36))
                ctk.CTkLabel(hdr, image=img, text="",
                             fg_color=PANEL, width=44, height=44).grid(
                                 row=0, column=0, padx=(20, 0), pady=10)
            except Exception:
                pass

        info = ctk.CTkFrame(hdr, fg_color="transparent")
        info.grid(row=0, column=1, sticky="w", padx=14)
        ctk.CTkLabel(info, text="Generador de Evaluaciones",
                     text_color=TEXT,
                     font=ctk.CTkFont(FONT, 15, "bold")).pack(anchor="w")
        ctk.CTkLabel(info, text="Con inteligencia artificial  ·  AMR  ·  by Levil",
                     text_color=MUTED,
                     font=ctk.CTkFont(FONT, 10)).pack(anchor="w")

        ctk.CTkFrame(self, fg_color=BORDER, height=1,
                     corner_radius=0).grid(row=0, column=0, sticky="sew")

    def _build_main(self):
        main = ctk.CTkFrame(self, fg_color=BG, corner_radius=0)
        main.grid(row=1, column=0, sticky="nsew", padx=14, pady=14)
        main.grid_columnconfigure(0, weight=0, minsize=210)
        main.grid_columnconfigure(1, weight=1)
        main.grid_columnconfigure(2, weight=0, minsize=272)
        main.grid_rowconfigure(0, weight=1)

        self._build_sidebar(main)
        self._build_chat_panel(main)
        self._build_opciones_panel(main)

    # ── Sidebar historial ─────────────────────────────────────────────────────
    def _build_sidebar(self, parent):
        side = ctk.CTkFrame(parent, fg_color=PANEL, corner_radius=16,
                            border_color=BORDER, border_width=1)
        side.grid(row=0, column=0, sticky="nsew", padx=(0, 10))
        side.grid_columnconfigure(0, weight=1)
        side.grid_rowconfigure(1, weight=1)

        hdr_side = ctk.CTkFrame(side, fg_color="transparent", height=46)
        hdr_side.grid(row=0, column=0, sticky="ew", padx=12, pady=(12, 0))
        hdr_side.grid_columnconfigure(0, weight=1)
        hdr_side.grid_propagate(False)

        ctk.CTkLabel(hdr_side, text="HISTORIAL",
                     text_color=MUTED,
                     font=ctk.CTkFont(FONT, 9, "bold")).grid(
                         row=0, column=0, sticky="w")

        ctk.CTkButton(hdr_side, text="+ Nueva",
                      fg_color=ACCENT, hover_color=ACCENT_D,
                      text_color="white",
                      font=ctk.CTkFont(FONT, 9, "bold"),
                      width=68, height=30, corner_radius=8,
                      command=self._nueva_conv).grid(row=0, column=1)

        ctk.CTkFrame(side, fg_color=BORDER, height=1,
                     corner_radius=0).grid(row=0, column=0, sticky="sew",
                                           padx=0, pady=(46, 0))

        self._sidebar_list = ctk.CTkScrollableFrame(
            side, fg_color="transparent",
            scrollbar_button_color=BORDER,
            scrollbar_button_hover_color=MUTED)
        self._sidebar_list.grid(row=1, column=0, sticky="nsew",
                                padx=6, pady=(6, 10))
        self._sidebar_list.grid_columnconfigure(0, weight=1)

    def _refrescar_sidebar(self):
        for w in self._sidebar_list.winfo_children():
            w.destroy()

        convs = _cargar_convs_disco()
        if not convs:
            ctk.CTkLabel(self._sidebar_list,
                         text="Sin conversaciones\nguardadas",
                         text_color=MUTED,
                         font=ctk.CTkFont(FONT, 10),
                         justify="center").pack(pady=24)
            return

        today     = date.today().isoformat()
        yesterday = date.fromordinal(date.today().toordinal() - 1).isoformat()
        semana    = date.fromordinal(date.today().toordinal() - 6).isoformat()

        grupos: dict = {}
        for conv in convs:
            d = conv.get("date", "")[:10]
            if d == today:
                g = "Hoy"
            elif d == yesterday:
                g = "Ayer"
            elif d >= semana:
                g = "Esta semana"
            else:
                g = "Anteriores"
            grupos.setdefault(g, []).append(conv)

        for nombre in ["Hoy", "Ayer", "Esta semana", "Anteriores"]:
            if nombre not in grupos:
                continue
            ctk.CTkLabel(self._sidebar_list, text=nombre,
                         text_color=MUTED,
                         font=ctk.CTkFont(FONT, 9, "bold"),
                         anchor="w").pack(fill="x", padx=8, pady=(10, 2))
            for conv in grupos[nombre]:
                self._sidebar_item(conv)

    def _sidebar_item(self, conv: dict):
        cid    = conv["id"]
        title  = conv.get("title", "Sin título")
        activo = (cid == self._conv_id)

        btn = ctk.CTkButton(
            self._sidebar_list,
            text=title,
            anchor="w",
            fg_color=ACCENT if activo else "transparent",
            hover_color=BORDER if not activo else ACCENT_D,
            text_color="white" if activo else TEXT2,
            font=ctk.CTkFont(FONT, 10),
            height=34, corner_radius=8,
            command=lambda c=cid: self._cargar_conv(c)
        )
        btn.pack(fill="x", padx=4, pady=2)
        btn.bind("<Button-3>", lambda e, c=cid, t=title: self._menu_conv(e, c, t))

    # ── Panel chat ────────────────────────────────────────────────────────────
    def _build_chat_panel(self, parent):
        chat_frame = ctk.CTkFrame(parent, fg_color=PANEL, corner_radius=16,
                                  border_color=BORDER, border_width=1)
        chat_frame.grid(row=0, column=1, sticky="nsew", padx=(0, 10))
        chat_frame.grid_columnconfigure(0, weight=1)
        chat_frame.grid_rowconfigure(1, weight=1)

        # Header
        hdr_chat = ctk.CTkFrame(chat_frame, fg_color=PANEL,
                                corner_radius=0, height=44)
        hdr_chat.grid(row=0, column=0, sticky="ew")
        hdr_chat.grid_propagate(False)
        ctk.CTkLabel(hdr_chat, text="CONVERSACIÓN",
                     text_color=MUTED,
                     font=ctk.CTkFont(FONT, 9, "bold")).pack(
                         side="left", padx=18, pady=12)
        ctk.CTkFrame(chat_frame, fg_color=BORDER, height=1,
                     corner_radius=0).grid(row=0, column=0, sticky="sew")

        # Área scrollable
        self._chat_area = ctk.CTkScrollableFrame(
            chat_frame, fg_color=CHAT_BG, corner_radius=0,
            scrollbar_button_color=BORDER,
            scrollbar_button_hover_color=MUTED)
        self._chat_area.grid(row=1, column=0, sticky="nsew")
        self._chat_area.grid_columnconfigure(0, weight=1)
        self.bind_all("<MouseWheel>", self._on_global_scroll)

        # Separador
        ctk.CTkFrame(chat_frame, fg_color=BORDER, height=1,
                     corner_radius=0).grid(row=2, column=0, sticky="ew")

        # Barra de input (sin altura fija — crece con el texto)
        input_bar = ctk.CTkFrame(chat_frame, fg_color=PANEL, corner_radius=0)
        input_bar.grid(row=3, column=0, sticky="ew", padx=14, pady=10)
        input_bar.grid_columnconfigure(2, weight=1)

        self._btn_arch = ctk.CTkButton(
            input_bar, text="📎",
            fg_color="transparent", hover_color=BG,
            text_color=MUTED, font=ctk.CTkFont(FONT, 15),
            width=38, height=52, corner_radius=8,
            command=self._subir_archivo)
        self._btn_arch.grid(row=0, column=0, padx=(0, 2))

        self._btn_img = ctk.CTkButton(
            input_bar, text="🖼",
            fg_color="transparent", hover_color=BG,
            text_color=MUTED, font=ctk.CTkFont(FONT, 14),
            width=38, height=52, corner_radius=8,
            command=self._subir_imagen)
        self._btn_img.grid(row=0, column=1, padx=(0, 6))

        self._chat_input = ctk.CTkTextbox(
            input_bar, height=52, wrap="word",
            font=ctk.CTkFont(FONT, 13),
            fg_color=BG, text_color=TEXT,
            border_color=BORDER, border_width=1,
            corner_radius=10)
        self._chat_input.grid(row=0, column=2, sticky="ew")
        self._chat_input.bind("<Return>",    self._on_enter)
        self._chat_input.bind("<KeyRelease>", self._ajustar_input)
        self._chat_input.bind("<Button-3>",  self._menu_contextual_input)

        self._btn_enviar = ctk.CTkButton(
            input_bar, text="→",
            fg_color=ACCENT, hover_color=ACCENT_D,
            text_color="white",
            font=ctk.CTkFont(FONT, 16, "bold"),
            width=48, height=52, corner_radius=10,
            command=self._enviar_chat)
        self._btn_enviar.grid(row=0, column=3, padx=(6, 0))

        self._lbl_arch = ctk.CTkLabel(chat_frame, text="",
                                      text_color=GREEN,
                                      font=ctk.CTkFont(FONT, 9))
        self._lbl_arch.grid(row=4, column=0, sticky="w", padx=18, pady=(0, 2))

        # Imagen preview (oculta hasta que se seleccione una)
        self._img_preview_frame = ctk.CTkFrame(chat_frame, fg_color="transparent")
        self._img_preview_frame.grid(row=5, column=0, sticky="w", padx=18, pady=(0, 4))
        self._img_preview_frame.grid_remove()

        self._lbl_img = ctk.CTkLabel(self._img_preview_frame, text="",
                                     text_color=PURPLE,
                                     font=ctk.CTkFont(FONT, 9))
        self._lbl_img.pack(side="left")
        ctk.CTkButton(self._img_preview_frame, text="✕",
                      fg_color="transparent", hover_color=BG,
                      text_color=MUTED, font=ctk.CTkFont(FONT, 9, "bold"),
                      width=22, height=18, corner_radius=4,
                      command=self._quitar_imagen).pack(side="left", padx=(6, 0))

    # ── Panel opciones ────────────────────────────────────────────────────────
    def _build_opciones_panel(self, parent):
        right = ctk.CTkScrollableFrame(
            parent, fg_color=PANEL, corner_radius=16,
            border_color=BORDER, border_width=1,
            scrollbar_button_color=BORDER,
            scrollbar_button_hover_color=MUTED)
        right.grid(row=0, column=2, sticky="nsew")
        right.grid_columnconfigure(0, weight=1)
        self._right_canvas = right._parent_canvas

        row = 0
        row = self._sec(right, "OPCIONES", row)

        opc = ctk.CTkFrame(right, fg_color="transparent")
        opc.grid(row=row, column=0, sticky="ew", padx=14, pady=(6, 0))
        opc.grid_columnconfigure(1, weight=1)
        row += 1

        opciones = [
            ("Asignatura",          ASIGNATURAS,                              "Lenguaje"),
            ("Curso",               CURSOS,                                   "3° Básico"),
            ("Tipo de preguntas",   TIPOS_PREGUNTA,                           "Selección múltiple"),
            ("Preguntas por ficha", [str(n) for n in range(2, 11)],           "4"),
            ("Cantidad de fichas",  [str(n) for n in range(1, 21)],           "1"),
            ("Tipo de documento",   ["Con encabezado", "Sin encabezado"],     "Con encabezado"),
        ]
        self._combos = []
        for i, (lbl, vals, default) in enumerate(opciones):
            ctk.CTkLabel(opc, text=lbl, text_color=MUTED,
                         font=ctk.CTkFont(FONT, 10)).grid(
                             row=i, column=0, sticky="w", pady=6)
            cb = ctk.CTkComboBox(opc, values=vals,
                                 font=ctk.CTkFont(FONT, 10),
                                 fg_color=BG, text_color=TEXT,
                                 border_color=BORDER, border_width=1,
                                 button_color=BORDER,
                                 button_hover_color=MUTED,
                                 dropdown_fg_color=PANEL,
                                 dropdown_text_color=TEXT,
                                 height=36, corner_radius=8,
                                 state="readonly")
            cb.set(default)
            cb.grid(row=i, column=1, sticky="ew", padx=(10, 0), pady=6)
            self._combos.append(cb)

        (self._cb_asig, self._cb_curso, self._cb_tipo_preg,
         self._cb_preg, self._cb_fich, self._cb_tipo) = self._combos

        row = self._sep(right, row)

        # ── Salida ────────────────────────────────────────────────────────────
        row = self._sec(right, "SALIDA WORD", row)

        self._seg_salida = ctk.CTkSegmentedButton(
            right,
            values=["Un solo Word", "Word por ficha"],
            font=ctk.CTkFont(FONT, 10),
            fg_color=BG,
            selected_color=ACCENT,
            selected_hover_color=ACCENT_D,
            unselected_color=BG,
            unselected_hover_color=BORDER,
            text_color=TEXT,
            text_color_disabled=MUTED,
            corner_radius=8, height=36)
        self._seg_salida.set("Un solo Word")
        self._seg_salida.grid(row=row, column=0, sticky="ew",
                              padx=14, pady=(6, 0))
        row += 1

        ctk.CTkLabel(right,
                     text="Word por ficha: máx. 10  ·  Un solo Word: sin límite",
                     text_color=MUTED,
                     font=ctk.CTkFont(FONT, 8)).grid(
                         row=row, column=0, sticky="w", padx=14, pady=(4, 0))
        row += 1

        row = self._sep(right, row)

        # ── Extras ────────────────────────────────────────────────────────────
        row = self._sec(right, "EXTRAS", row)

        self._var_clave   = tk.BooleanVar()
        self._var_version = tk.BooleanVar()

        chk_frame = ctk.CTkFrame(right, fg_color="transparent")
        chk_frame.grid(row=row, column=0, sticky="ew", padx=14, pady=(6, 0))
        row += 1

        for var, texto in [
            (self._var_clave,   "Generar clave de respuestas"),
            (self._var_version, "Generar Versión A y B"),
        ]:
            ctk.CTkCheckBox(chk_frame, text=texto, variable=var,
                            font=ctk.CTkFont(FONT, 10),
                            text_color=TEXT2,
                            fg_color=ACCENT,
                            hover_color=ACCENT_D,
                            border_color=BORDER,
                            checkmark_color="white",
                            corner_radius=5).pack(anchor="w", pady=5)

        row = self._sep(right, row)

        # ── Botones ───────────────────────────────────────────────────────────
        btns = ctk.CTkFrame(right, fg_color="transparent")
        btns.grid(row=row, column=0, sticky="ew", padx=14, pady=(10, 16))
        btns.grid_columnconfigure(0, weight=1)
        row += 1

        self._btn_dl = ctk.CTkButton(
            btns, text="Descargar Word",
            fg_color=GREEN, hover_color=GREEN_D,
            text_color="white",
            text_color_disabled="#6B7280",
            font=ctk.CTkFont(FONT, 13, "bold"),
            height=48, corner_radius=12,
            state="disabled",
            command=self._descargar)
        self._btn_dl.grid(row=0, column=0, sticky="ew", pady=(0, 8))

        ctk.CTkButton(
            btns, text="Cambiar API Key",
            fg_color=BORDER, hover_color="#3A3A3C",
            text_color=TEXT2,
            font=ctk.CTkFont(FONT, 10),
            height=32, corner_radius=8,
            command=self._cambiar_api_key).grid(row=1, column=0, sticky="ew")

    # ── Helpers de layout ─────────────────────────────────────────────────────
    def _sec(self, parent, texto, row):
        ctk.CTkLabel(parent, text=texto,
                     text_color=MUTED,
                     font=ctk.CTkFont(FONT, 9, "bold")).grid(
                         row=row, column=0, sticky="w",
                         padx=14, pady=(18, 0))
        return row + 1

    def _sep(self, parent, row):
        ctk.CTkFrame(parent, fg_color=BORDER,
                     height=1, corner_radius=0).grid(
                         row=row, column=0, sticky="ew",
                         padx=14, pady=(10, 0))
        return row + 1

    # ── Resize ────────────────────────────────────────────────────────────────
    def _on_resize(self, event):
        if event.widget is not self:
            return
        if self._resize_job:
            self.after_cancel(self._resize_job)
        self._resize_job = self.after(120, self._aplicar_resize)

    def _aplicar_resize(self):
        self._resize_job = None
        w = self._chat_area.winfo_width()
        if w > 100:
            self._chat_w = w

    # ── Burbujas ──────────────────────────────────────────────────────────────
    def _burbuja(self, texto: str, rol: str, guardar: bool = True):
        es_user = (rol == "user")
        row     = len(self._chat_area.winfo_children())

        outer = ctk.CTkFrame(self._chat_area, fg_color="transparent")
        outer.grid(row=row, column=0, sticky="ew", padx=14, pady=(6, 2))
        outer.grid_columnconfigure(0, weight=1)

        now = datetime.now().strftime("%H:%M")

        if es_user:
            ancho   = max(220, int(self._chat_w * 0.63))
            content = ctk.CTkFrame(outer, fg_color="transparent")
            content.grid(row=0, column=0, sticky="e")

            bub = ctk.CTkFrame(content, fg_color=BUB_U, corner_radius=20)
            bub.pack(anchor="e")
            ctk.CTkLabel(bub, text=texto,
                         text_color="white",
                         wraplength=ancho, justify="left",
                         font=ctk.CTkFont(FONT, 11),
                         padx=16, pady=12).pack()

            ctk.CTkLabel(content, text=now,
                         text_color=MUTED,
                         font=ctk.CTkFont(FONT, 8)).pack(anchor="e", padx=4, pady=(3, 0))
        else:
            ancho = max(220, int(self._chat_w * 0.76))

            bubble_row = ctk.CTkFrame(outer, fg_color="transparent")
            bubble_row.grid(row=0, column=0, sticky="w")

            # Ícono circular "G"
            icon_wrap = ctk.CTkFrame(bubble_row, fg_color=ACCENT,
                                     width=28, height=28, corner_radius=14)
            icon_wrap.pack(side="left", anchor="n", padx=(0, 8), pady=4)
            icon_wrap.pack_propagate(False)
            ctk.CTkLabel(icon_wrap, text="G", text_color="white",
                         font=ctk.CTkFont(FONT, 10, "bold")).place(
                             relx=0.5, rely=0.5, anchor="center")

            content = ctk.CTkFrame(bubble_row, fg_color="transparent")
            content.pack(side="left", anchor="n")

            bub = ctk.CTkFrame(content, fg_color=BUB_IA,
                               border_color=BORDER, border_width=1,
                               corner_radius=20)
            bub.pack(anchor="w")
            ctk.CTkLabel(bub, text=texto,
                         text_color=TEXT2,
                         wraplength=ancho, justify="left",
                         font=ctk.CTkFont(FONT, 11),
                         padx=16, pady=12).pack()

            ctk.CTkLabel(content, text=now,
                         text_color=MUTED,
                         font=ctk.CTkFont(FONT, 8)).pack(anchor="w", padx=4, pady=(3, 0))

        self.update_idletasks()
        self._chat_area._parent_canvas.yview_moveto(1.0)

        if guardar:
            self._conv_messages.append({"rol": rol, "texto": texto})

        return outer

    def _burbuja_pensando(self):
        w = self._burbuja("Pensando…", "ia", guardar=False)
        self._pensando_widget = w

    def _quitar_pensando(self):
        if self._pensando_widget:
            self._pensando_widget.destroy()
            self._pensando_widget = None

    # ── Menú contextual ───────────────────────────────────────────────────────
    def _menu_contextual_input(self, event):
        menu = tk.Menu(self, tearoff=0,
                       bg=PANEL, fg=TEXT,
                       activebackground=ACCENT, activeforeground="white",
                       font=(FONT, 10), relief="flat", bd=0)
        menu.add_command(label="Cortar",
                         command=lambda: self._chat_input.event_generate("<<Cut>>"))
        menu.add_command(label="Copiar",
                         command=lambda: self._chat_input.event_generate("<<Copy>>"))
        menu.add_command(label="Pegar",
                         command=lambda: self._chat_input.event_generate("<<Paste>>"))
        menu.add_separator()
        menu.add_command(label="Seleccionar todo",
                         command=lambda: self._chat_input.tag_add("sel", "1.0", "end"))
        try:
            menu.tk_popup(event.x_root, event.y_root)
        finally:
            menu.grab_release()

    # ── Scroll global ─────────────────────────────────────────────────────────
    def _on_global_scroll(self, event):
        delta = int(-1 * (event.delta / 120)) * 3
        mx, my = event.x_root, event.y_root

        def _sobre(canvas):
            try:
                x = canvas.winfo_rootx()
                y = canvas.winfo_rooty()
                return x <= mx <= x + canvas.winfo_width() and \
                       y <= my <= y + canvas.winfo_height()
            except Exception:
                return False

        if _sobre(self._chat_area._parent_canvas):
            self._chat_area._parent_canvas.yview_scroll(delta, "units")
            return "break"
        if _sobre(self._right_canvas):
            self._right_canvas.yview_scroll(delta, "units")
            return "break"

    # ── Adjuntar archivo ──────────────────────────────────────────────────────
    def _subir_archivo(self):
        path = filedialog.askopenfilename(
            title="Seleccionar archivo",
            filetypes=[("Soportados", "*.pdf *.docx *.txt"),
                       ("PDF", "*.pdf"), ("Word", "*.docx"), ("Texto", "*.txt")])
        if not path:
            return
        try:
            self._texto_base     = extraer_texto(path)
            self._archivo_nombre = os.path.basename(path)
            self._lbl_arch.configure(text=f"  Adjunto: {self._archivo_nombre}")
        except Exception as e:
            messagebox.showerror("Error al leer archivo", str(e))
            return
        if path.lower().endswith(".pdf"):
            self._intentar_extraer_imagenes_pdf(path)

    def _subir_imagen(self):
        """Selector manual de imagen JPG/PNG."""
        path = filedialog.askopenfilename(
            title="Seleccionar imagen",
            filetypes=[("Imágenes", "*.jpg *.jpeg *.png *.bmp"),
                       ("JPEG", "*.jpg *.jpeg"), ("PNG", "*.png")])
        if not path:
            return
        try:
            with open(path, "rb") as f:
                self._seleccionar_imagen(f.read(), os.path.basename(path))
        except Exception as e:
            messagebox.showerror("Error al leer imagen", str(e))

    def _intentar_extraer_imagenes_pdf(self, path: str):
        """Extrae imágenes del PDF con pymupdf y abre el picker si hay varias."""
        try:
            import fitz
        except ImportError:
            return
        try:
            doc = fitz.open(path)
            imagenes = []
            for page_num in range(len(doc)):
                for img_info in doc[page_num].get_images(full=True):
                    xref = img_info[0]
                    base_img = doc.extract_image(xref)
                    img_bytes = base_img["image"]
                    if len(img_bytes) > 5000:  # omitir íconos tiny
                        ext = base_img["ext"]
                        imagenes.append({
                            "bytes":  img_bytes,
                            "ext":    ext,
                            "nombre": f"img_pag{page_num+1}_{len(imagenes)+1}.{ext}",
                        })
            doc.close()
            if len(imagenes) == 1:
                self._seleccionar_imagen(imagenes[0]["bytes"], imagenes[0]["nombre"])
            elif len(imagenes) > 1:
                self._mostrar_picker_imagenes(imagenes)
        except Exception:
            pass

    def _mostrar_picker_imagenes(self, imagenes: list):
        """Diálogo con miniaturas para elegir qué imagen incluir."""
        from PIL import Image
        from io import BytesIO

        dlg = ctk.CTkToplevel(self)
        dlg.title("Elegir imagen del PDF")
        dlg.geometry("520x400")
        dlg.configure(fg_color=PANEL)
        dlg.grab_set()

        ctk.CTkLabel(dlg, text="Elige una imagen para incluir en la evaluación",
                     text_color=TEXT,
                     font=ctk.CTkFont(FONT, 12, "bold")).pack(padx=20, pady=(20, 10))

        scroll = ctk.CTkScrollableFrame(dlg, fg_color=BG,
                                        scrollbar_button_color=BORDER)
        scroll.pack(fill="both", expand=True, padx=16, pady=(0, 6))

        selected = [None]

        def _pick(img_data):
            selected[0] = img_data
            dlg.destroy()

        for i, img_data in enumerate(imagenes):
            try:
                pil_img = Image.open(BytesIO(img_data["bytes"]))
                pil_img.thumbnail((120, 120))
                ctk_img = ctk.CTkImage(pil_img, size=(pil_img.width, pil_img.height))
            except Exception:
                ctk_img = None

            btn = ctk.CTkButton(
                scroll,
                text=img_data["nombre"],
                image=ctk_img,
                compound="top",
                fg_color=BG, hover_color=BORDER,
                text_color=TEXT2,
                font=ctk.CTkFont(FONT, 8),
                width=140, height=150,
                command=lambda d=img_data: _pick(d))
            btn.grid(row=i // 3, column=i % 3, padx=8, pady=8)

        ctk.CTkButton(dlg, text="Sin imagen",
                      fg_color="transparent", hover_color=BORDER,
                      text_color=MUTED, font=ctk.CTkFont(FONT, 10),
                      command=dlg.destroy).pack(pady=(0, 14))

        self.wait_window(dlg)
        if selected[0]:
            self._seleccionar_imagen(selected[0]["bytes"], selected[0]["nombre"])

    def _seleccionar_imagen(self, img_bytes: bytes, nombre: str):
        self._imagen_bytes  = img_bytes
        self._imagen_nombre = nombre
        self._lbl_img.configure(text=f"  🖼 {nombre}")
        self._img_preview_frame.grid()

    def _quitar_imagen(self):
        self._imagen_bytes  = None
        self._imagen_nombre = ""
        self._img_preview_frame.grid_remove()

    # ── Chat ──────────────────────────────────────────────────────────────────
    def _ajustar_input(self, _=None):
        """Auto-expande el input según el contenido, máximo 3 líneas."""
        try:
            lines = int(self._chat_input.count("1.0", "end-1c", "displaylines")[0])
        except Exception:
            lines = 1
        lines = max(1, min(3, lines))
        new_h = 52 + (lines - 1) * 24
        self._chat_input.configure(height=new_h)

    def _on_enter(self, event):
        # Shift+Enter = nueva línea, Enter solo = enviar
        if event.state & 0x1:
            return
        self._enviar_chat()
        return "break"

    def _enviar_chat(self):
        if self._generando:
            return
        if not self._gemini:
            messagebox.showerror("Sin API Key", "Configura tu API Key primero.")
            return

        prompt = self._chat_input.get("1.0", "end").strip()
        if not prompt:
            return

        self._chat_input.delete("1.0", "end")
        self._chat_input.configure(height=52)

        n_fichas      = int(self._cb_fich.get())
        n_preguntas   = int(self._cb_preg.get())
        curso         = self._cb_curso.get()
        asignatura    = self._cb_asig.get()
        tipo_pregunta = self._cb_tipo_preg.get()

        if self._seg_salida.get() == "Word por ficha" and n_fichas > 10:
            messagebox.showwarning("Límite excedido",
                "En modo «Word por ficha» el máximo son 10 fichas.\n"
                "Reduce la cantidad o cambia a «Un solo Word».")
            return

        self._generando = True
        self._btn_enviar.configure(state="disabled", text="…")
        self._btn_dl.configure(state="disabled")

        meta = (f"{n_fichas} ficha(s)  ·  {n_preguntas} preguntas  "
                f"·  {tipo_pregunta}  ·  {asignatura}  ·  {curso}")
        if self._archivo_nombre:
            meta += f"  ·  {self._archivo_nombre}"
        self._burbuja(f"{prompt}\n\n{meta}", "user")
        self._burbuja_pensando()

        threading.Thread(
            target=self._worker,
            args=(prompt, n_fichas, n_preguntas, curso, asignatura, tipo_pregunta,
                  self._imagen_bytes),
            daemon=True
        ).start()

    def _worker(self, prompt, n_fichas, n_preguntas, curso, asignatura, tipo_pregunta,
                imagen_bytes=None):
        try:
            fichas = self._gemini.generar(
                prompt=prompt, n_fichas=n_fichas,
                n_preguntas=n_preguntas, curso=curso,
                asignatura=asignatura,
                tipo_pregunta=tipo_pregunta,
                texto_base=self._texto_base,
                imagen_bytes=imagen_bytes)
            self.after(0, self._on_ok, fichas)
        except Exception as e:
            self.after(0, self._on_err, str(e))

    def _on_ok(self, fichas):
        self._fichas    = fichas
        self._generando = False
        self._quitar_pensando()
        self._btn_enviar.configure(state="normal", text="→")
        self._btn_dl.configure(state="normal")
        total = sum(len(f.get("preguntas", [])) for f in fichas)
        msg = (f"Listo. Generé {len(fichas)} ficha(s) con {total} pregunta(s) en total.\n"
               "Pulsa «Descargar Word» cuando quieras.")
        self._burbuja(msg, "ia")
        self._auto_guardar_conv()

    def _on_err(self, err):
        self._generando = False
        self._quitar_pensando()
        self._btn_enviar.configure(state="normal", text="→")
        self._burbuja(f"Ocurrió un error:\n{err}", "ia")

    # ── Persistencia ──────────────────────────────────────────────────────────
    def _auto_guardar_conv(self):
        if not self._conv_id:
            self._conv_id = datetime.now().strftime("%Y%m%d_%H%M%S_") + \
                            str(uuid.uuid4())[:6]

        titulo = "Sin título"
        for m in self._conv_messages:
            if m["rol"] == "user":
                t = m["texto"].split("\n")[0][:40].strip()
                titulo = t + ("…" if len(m["texto"]) > 40 else "")
                break

        conv = {
            "id":             self._conv_id,
            "title":          titulo,
            "date":           datetime.now().isoformat(),
            "messages":       self._conv_messages.copy(),
            "gemini_history": self._gemini.get_historial_raw() if self._gemini else [],
            "fichas":         self._fichas or [],
        }
        _guardar_conv_disco(conv)
        self._refrescar_sidebar()

    def _menu_conv(self, event, conv_id: str, titulo: str):
        menu = tk.Menu(self, tearoff=0,
                       bg=PANEL, fg=TEXT,
                       activebackground=ACCENT, activeforeground="white",
                       font=(FONT, 10), relief="flat", bd=0)
        menu.add_command(label="Renombrar",
                         command=lambda: self._renombrar_conv(conv_id, titulo))
        menu.add_separator()
        menu.add_command(label="Eliminar",
                         command=lambda: self._eliminar_conv(conv_id))
        try:
            menu.tk_popup(event.x_root, event.y_root)
        finally:
            menu.grab_release()

    def _renombrar_conv(self, conv_id: str, titulo_actual: str):
        dlg = ctk.CTkToplevel(self)
        dlg.title("Renombrar conversación")
        dlg.geometry("400x180")
        dlg.resizable(False, False)
        dlg.configure(fg_color=PANEL)
        dlg.grab_set()

        ctk.CTkLabel(dlg, text="Nuevo nombre",
                     text_color=TEXT,
                     font=ctk.CTkFont(FONT, 14, "bold")).pack(
                         padx=28, pady=(24, 8), anchor="w")

        entry = ctk.CTkEntry(dlg, width=344,
                             font=ctk.CTkFont(FONT, 12),
                             fg_color=BG, text_color=TEXT,
                             border_color=BORDER, border_width=1,
                             height=40, corner_radius=8)
        entry.pack(padx=28)
        entry.insert(0, titulo_actual)
        entry.select_range(0, "end")
        entry.focus()

        def _guardar(e=None):
            nuevo = entry.get().strip()
            if not nuevo:
                return
            path = os.path.join(_CONV_DIR, f"{conv_id}.json")
            if os.path.isfile(path):
                try:
                    with open(path, encoding="utf-8") as f:
                        conv = json.load(f)
                    conv["title"] = nuevo
                    with open(path, "w", encoding="utf-8") as f:
                        json.dump(conv, f, ensure_ascii=False, indent=2)
                except Exception:
                    pass
            dlg.destroy()
            self._refrescar_sidebar()

        entry.bind("<Return>", _guardar)
        ctk.CTkButton(dlg, text="Guardar",
                      fg_color=ACCENT, hover_color=ACCENT_D,
                      text_color="white",
                      font=ctk.CTkFont(FONT, 11, "bold"),
                      height=38, corner_radius=8,
                      command=_guardar).pack(padx=28, pady=(12, 0), fill="x")

        self.wait_window(dlg)

    def _eliminar_conv(self, conv_id: str):
        if not messagebox.askyesno("Eliminar conversación",
                                   "¿Eliminar esta conversación? No se puede deshacer."):
            return
        path = os.path.join(_CONV_DIR, f"{conv_id}.json")
        try:
            os.remove(path)
        except Exception:
            pass
        if conv_id == self._conv_id:
            self._nueva_conv(silencioso=True)
        else:
            self._refrescar_sidebar()

    def _cargar_conv(self, conv_id: str):
        if self._conv_messages:
            self._auto_guardar_conv()

        path = os.path.join(_CONV_DIR, f"{conv_id}.json")
        if not os.path.isfile(path):
            return
        try:
            with open(path, encoding="utf-8") as f:
                conv = json.load(f)
        except Exception:
            return

        self._conv_id       = conv["id"]
        self._conv_messages = conv.get("messages", [])
        self._fichas        = conv.get("fichas") or None

        if self._gemini:
            self._gemini.set_historial_raw(conv.get("gemini_history", []))

        for w in self._chat_area.winfo_children():
            w.destroy()
        for m in self._conv_messages:
            self._burbuja(m["texto"], m["rol"], guardar=False)

        if self._fichas:
            self._btn_dl.configure(state="normal")
        else:
            self._btn_dl.configure(state="disabled")

        self._refrescar_sidebar()

    def _nueva_conv(self, silencioso: bool = False):
        if not silencioso and self._conv_messages:
            self._auto_guardar_conv()

        self._conv_id        = None
        self._conv_messages  = []
        self._fichas         = None
        self._texto_base     = ""
        self._archivo_nombre = ""
        self._imagen_bytes   = None
        self._imagen_nombre  = ""
        self._pensando_widget = None

        if self._gemini:
            self._gemini.limpiar_historial()

        self._lbl_arch.configure(text="")
        self._btn_dl.configure(state="disabled")
        if hasattr(self, "_img_preview_frame"):
            self._img_preview_frame.grid_remove()

        for w in self._chat_area.winfo_children():
            w.destroy()

        self._burbuja(
            "Hola. Describe la evaluación que necesitas y la genero al instante.\n\n"
            "Ejemplo: «Prueba de comprensión sobre el agua para 3° básico, 4 preguntas»",
            "ia", guardar=False)

        self._refrescar_sidebar()

    # ── Descargar ─────────────────────────────────────────────────────────────
    def _descargar(self):
        if not self._fichas:
            return

        ts     = datetime.now().strftime("%Y%m%d_%H%M%S")
        layout = "doble" if self._cb_tipo.get() == "Con encabezado" else "normal"
        modo   = self._seg_salida.get()
        hechos = []

        try:
            if modo == "Word por ficha":
                fichas_out = self._fichas[:10]
                folder = filedialog.askdirectory(
                    title="Seleccionar carpeta donde guardar las fichas")
                if not folder:
                    return

                for i, ficha in enumerate(fichas_out, 1):
                    if self._var_version.get():
                        path_a = os.path.join(folder,
                                              f"ficha{i}_versionA_{ts}.docx")
                        Generador().generar([ficha], path_a, layout,
                                            imagen_bytes=self._imagen_bytes)
                        hechos.append(f"Ficha {i} — Versión A")

                        fb = _hacer_version_b(ficha)
                        path_b = os.path.join(folder,
                                              f"ficha{i}_versionB_{ts}.docx")
                        Generador().generar([fb], path_b, layout,
                                            imagen_bytes=self._imagen_bytes)
                        hechos.append(f"Ficha {i} — Versión B")
                    else:
                        path = os.path.join(folder, f"ficha{i}_{ts}.docx")
                        Generador().generar([ficha], path, layout,
                                            imagen_bytes=self._imagen_bytes)
                        hechos.append(f"Ficha {i}")

                if self._var_clave.get():
                    path_clave = os.path.join(folder, f"clave_{ts}.docx")
                    generar_clave_docx(fichas_out, path_clave)
                    hechos.append("Clave de respuestas")

            else:  # Un solo Word
                out = filedialog.asksaveasfilename(
                    title="Guardar evaluación",
                    initialfile=f"evaluacion_{ts}.docx",
                    defaultextension=".docx",
                    filetypes=[("Word", "*.docx")])
                if not out:
                    return
                base = out[:-5]

                if self._var_version.get():
                    Generador().generar(self._fichas,
                                        f"{base}_versionA.docx", layout,
                                        imagen_bytes=self._imagen_bytes)
                    hechos.append("Versión A")
                    fb_list = [_hacer_version_b(f) for f in self._fichas]
                    Generador().generar(fb_list, f"{base}_versionB.docx", layout,
                                        imagen_bytes=self._imagen_bytes)
                    hechos.append("Versión B")
                else:
                    Generador().generar(self._fichas, out, layout,
                                        imagen_bytes=self._imagen_bytes)
                    hechos.append("Evaluación")

                if self._var_clave.get():
                    generar_clave_docx(self._fichas, f"{base}_respuestas.docx")
                    hechos.append("Clave de respuestas")

        except Exception as e:
            messagebox.showerror("Error al guardar", str(e))
            return

        msg = "Guardado:\n" + "\n".join(f"  · {h}" for h in hechos)
        self._burbuja(msg, "ia")
        messagebox.showinfo("Listo", msg)

    def _cambiar_api_key(self):
        self._pedir_api_key()


# ════════════════════════════════════════════════════════════════════════════
if __name__ == "__main__":
    App().mainloop()
